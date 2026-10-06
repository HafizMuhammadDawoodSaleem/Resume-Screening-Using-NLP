"""Fast local tests. No model download or internet needed."""
from io import BytesIO, StringIO
import unittest

import numpy as np
from docx import Document

from resume_screening.documents import Resume, clean_text, extract_text, load_resumes
from resume_screening.ranking import KeywordEncoder, SemanticEncoder, rank_resumes


class FakeEncoder:
    def chunk(self, text):
        return [text]

    def encode(self, texts):
        return np.array([[1., 0.] if "python" in text.lower() else [0., 1.] for text in texts])


class ProjectTests(unittest.TestCase):
    def test_cosine_order_and_explanation(self):
        matches = rank_resumes("Python developer", [Resume("1", "Designer", "Graphic design."),
                               Resume("2", "Engineer", "Python developer.")], FakeEncoder())
        self.assertEqual([m.resume_id for m in matches], ["2", "1"])
        self.assertAlmostEqual(matches[0].cosine_similarity, 1.)
        self.assertAlmostEqual(matches[1].cosine_similarity, 0.)
        self.assertEqual(matches[0].evidence, "Python developer.")

    def test_tfidf_distinguishes_relevant_resume(self):
        matches = rank_resumes("Python SQL data analysis", [Resume("a", "A", "Python SQL data analysis"),
                                Resume("b", "B", "Nursing patient care")], KeywordEncoder(), 1)
        self.assertEqual(matches[0].resume_id, "a")

    def test_long_resume_tail_is_compared(self):
        resume = "Nursing patient care. " * 150 + "Python SQL data analysis."
        match = rank_resumes("Python SQL data analysis", [Resume("a", "A", resume)], KeywordEncoder())[0]
        self.assertIn("Python SQL", match.evidence)

    def test_stable_ties_and_duplicate_ids(self):
        resumes = [Resume("a", "A", "Python"), Resume("b", "B", "Python")]
        self.assertEqual([m.resume_id for m in rank_resumes("Python", resumes, FakeEncoder())], ["a", "b"])
        with self.assertRaises(ValueError):
            rank_resumes("Python", [resumes[0], resumes[0]], FakeEncoder())

    def test_empty_and_invalid_inputs(self):
        for job, resumes, top in [("", [Resume("a", "A", "Python")], 1),
                                   ("Python", [], 1), ("Python", [Resume("a", "A", "Python")], 0)]:
            with self.assertRaises(ValueError):
                rank_resumes(job, resumes, FakeEncoder(), top)
        for name, content in [("empty.txt", b""), ("scan.txt", b"   "), ("old.doc", b"x"),
                               ("bad.pdf", b"invalid"), ("bad.txt", b"\xff")]:
            with self.assertRaises(ValueError):
                extract_text(name, content)

    def test_text_and_docx_table_extraction(self):
        self.assertEqual(extract_text("resume.txt", b"Python\n SQL"), "Python SQL")
        document = Document()
        document.add_paragraph("Python engineer")
        document.add_table(rows=1, cols=1).cell(0, 0).text = "SQL experience"
        buffer = BytesIO()
        document.save(buffer)
        self.assertIn("SQL experience", extract_text("resume.docx", buffer.getvalue()))

    def test_csv_validation_and_mapping(self):
        csv = StringIO("ID,Category,Resume_str\n1,Engineer,Python SQL\n")
        self.assertEqual(load_resumes(csv, "Resume_str", "ID", "Category")[0].text, "Python SQL")
        with self.assertRaises(ValueError):
            load_resumes(StringIO("id,name,resume_text\n1,A,Python\n1,B,SQL\n"))
        with self.assertRaises(ValueError):
            load_resumes(StringIO("id,name,resume_text\n1,A,\n"))

    def test_preprocessing_preserves_technical_terms_and_negation(self):
        self.assertEqual(clean_text("C++  C#\nnot Java"), "C++ C# not Java")

    def test_semantic_chunking_uses_token_limit_and_keeps_tail(self):
        class Tokenizer:
            def encode(self, text, **kwargs):
                return list(range(600))
            def decode(self, tokens, **kwargs):
                return " ".join(map(str, tokens))
        class Model:
            tokenizer = Tokenizer()
            max_seq_length = 256
        encoder = object.__new__(SemanticEncoder)
        encoder.model = Model()
        chunks = encoder.chunk("a long resume")
        self.assertTrue(all(len(chunk.split()) <= 240 for chunk in chunks))
        self.assertEqual(chunks[-1].split()[-1], "599")


if __name__ == "__main__":
    unittest.main()
