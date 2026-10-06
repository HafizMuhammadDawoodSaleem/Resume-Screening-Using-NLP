"""Read local documents without saving uploaded content."""
from dataclasses import dataclass
from io import BytesIO
from pathlib import Path
import re
import unicodedata
import zipfile

import pandas as pd

MAX_BYTES = 5 * 1024 * 1024
MAX_CHARS = 100_000


@dataclass(frozen=True)
class Resume:
    id: str
    name: str
    text: str


def clean_text(text: str) -> str:
    """Preserve punctuation, case and negation for contextual embeddings."""
    text = unicodedata.normalize("NFKC", str(text)).replace("\x00", " ")
    return re.sub(r"\s+", " ", text).strip()


def validate_text(text: str) -> str:
    text = clean_text(text)
    if not text or not re.search(r"\w", text):
        raise ValueError("No readable text found. Scanned PDFs require OCR first.")
    if len(text) > MAX_CHARS:
        raise ValueError(f"Text exceeds {MAX_CHARS:,} characters; shorten the document.")
    return text


def extract_text(filename: str, content: bytes) -> str:
    if not content:
        raise ValueError("The file is empty.")
    if len(content) > MAX_BYTES:
        raise ValueError("Maximum file size is 5 MB.")
    suffix = Path(filename).suffix.lower()
    try:
        if suffix == ".txt":
            text = content.decode("utf-8-sig")
        elif suffix == ".pdf":
            from pypdf import PdfReader
            reader = PdfReader(BytesIO(content))
            if reader.is_encrypted:
                raise ValueError("Password-protected PDFs are not supported.")
            if len(reader.pages) > 50:
                raise ValueError("PDFs must contain no more than 50 pages.")
            text = "\n".join(page.extract_text() or "" for page in reader.pages)
        elif suffix == ".docx":
            from docx import Document
            with zipfile.ZipFile(BytesIO(content)) as archive:
                if sum(info.file_size for info in archive.infolist()) > 25 * 1024 * 1024:
                    raise ValueError("DOCX decompressed contents exceed 25 MB.")
            doc = Document(BytesIO(content))
            parts = [paragraph.text for paragraph in doc.paragraphs]
            parts += [cell.text for table in doc.tables for row in table.rows for cell in row.cells]
            text = "\n".join(parts)
        else:
            raise ValueError("Supported files: .txt, .pdf and .docx.")
    except UnicodeDecodeError as exc:
        raise ValueError("Save the TXT file with UTF-8 encoding.") from exc
    except ValueError:
        raise
    except Exception as exc:
        raise ValueError("Could not read the file. Check that it is valid and not encrypted.") from exc
    return validate_text(text)


def read_csv(path, text_column: str, id_column: str, name_column: str) -> pd.DataFrame:
    frame = pd.read_csv(path, dtype=str, keep_default_na=False)
    required = {text_column, id_column, name_column}
    missing = required - set(frame.columns)
    if missing:
        raise ValueError(f"Missing CSV columns: {', '.join(sorted(missing))}")
    if frame.empty:
        raise ValueError("The CSV contains no records.")
    frame = frame.rename(columns={text_column: "text", id_column: "id", name_column: "name"})
    frame = frame[["id", "name", "text"]].copy()
    for column in frame:
        frame[column] = frame[column].map(clean_text)
        if frame[column].eq("").any():
            raise ValueError(f"CSV contains blank {column} values.")
    if frame.id.duplicated().any():
        raise ValueError("CSV IDs must be unique.")
    frame["text"] = frame.text.map(validate_text)
    return frame


def load_resumes(path, text_column="resume_text", id_column="id", name_column="name"):
    return [Resume(**row) for row in read_csv(path, text_column, id_column, name_column).to_dict("records")]


def load_jobs(path, text_column="description", id_column="id", name_column="title"):
    return read_csv(path, text_column, id_column, name_column)
